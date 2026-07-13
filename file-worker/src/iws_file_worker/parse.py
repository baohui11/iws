from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl
from docx import Document
from docx.document import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader

from .config import Settings
from .models import FileRow
from .paddleocr import PaddleOcrClient
from .preview import convert_office_to_pdf, decode_text, normalize_ext
from .storage import ObjectStorage

TEXT_EXT = {"txt", "log"}
MARKDOWN_EXT = {"md", "markdown"}
CSV_EXT = {"csv", "tsv"}
DOCX_EXT = {"docx"}
XLSX_EXT = {"xlsx"}
PDF_EXT = {"pdf"}
PPT_EXT = {"ppt", "pptx"}
OLD_OFFICE_EXT = {"doc", "xls"}

MAX_CHUNK_CHARS = 1024
CHUNK_OVERLAP_CHARS = 50
TEXT_DENSITY_MIN_CHARS_PER_PAGE = 30
SLIDE_SCORE_THRESHOLD = 3
EXCEL_WINDOW_ROWS = 80


@dataclass
class ParsedChunk:
    chunk_index: int
    content: str
    page_no: int | None = None
    slide_no: int | None = None
    sheet_name: str | None = None
    row_start: int | None = None
    row_end: int | None = None
    section_title: str | None = None
    metadata: dict[str, Any] = field(default_factory=dict)


@dataclass
class ParseResult:
    status: str
    storage_key: str | None = None
    content_text: str = ""
    chunks: list[ParsedChunk] = field(default_factory=list)
    parser_name: str = "iws-file-worker-parser"
    parser_version: str = "0.1.0"
    language: str | None = "zh"
    metadata: dict[str, Any] = field(default_factory=dict)
    output: dict[str, Any] | None = None
    error_code: str | None = None
    error_message: str | None = None


@dataclass
class PdfProfile:
    kind: str
    page_count: int
    text_len: int
    avg_chars_per_page: float
    landscape_ratio: float
    short_line_ratio: float
    avg_line_len: float
    slide_score: int
    metadata: dict[str, Any] = field(default_factory=dict)


def with_chunk_overlap(chunks: list[str], overlap_chars: int = CHUNK_OVERLAP_CHARS) -> list[str]:
    if overlap_chars <= 0 or len(chunks) <= 1:
        return chunks
    overlapped = [chunks[0]]
    for previous, current in zip(chunks, chunks[1:]):
        overlap = previous[-overlap_chars:].strip()
        overlapped.append(f"{overlap}\n{current}".strip() if overlap else current)
    return overlapped


def split_long_text(
    text: str,
    max_chars: int = MAX_CHUNK_CHARS,
    overlap_chars: int = CHUNK_OVERLAP_CHARS,
) -> list[str]:
    normalized = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not normalized:
        return []
    parts = re.split(r"\n\s*\n", normalized)
    chunks: list[str] = []
    buf = ""
    for part in parts:
        p = part.strip()
        if not p:
            continue
        if len(p) > max_chars:
            if buf:
                chunks.append(buf.strip())
                buf = ""
            for i in range(0, len(p), max_chars):
                chunks.append(p[i : i + max_chars].strip())
            continue
        if len(buf) + len(p) + 2 > max_chars:
            if buf:
                chunks.append(buf.strip())
            buf = p
        else:
            buf = f"{buf}\n\n{p}".strip() if buf else p
    if buf:
        chunks.append(buf.strip())
    return with_chunk_overlap(chunks, overlap_chars)


def row_to_text(row: list[Any]) -> str:
    cells = ["" if v is None else str(v).strip() for v in row]
    while cells and not cells[-1]:
        cells.pop()
    return "\t".join(cells)


def append_text_chunks(
    chunks: list[ParsedChunk],
    text: str,
    *,
    section_title: str | None = None,
    metadata: dict[str, Any] | None = None,
    page_no: int | None = None,
    slide_no: int | None = None,
    sheet_name: str | None = None,
    row_start: int | None = None,
    row_end: int | None = None,
) -> None:
    for content in split_long_text(text):
        chunks.append(
            ParsedChunk(
                len(chunks),
                content,
                page_no=page_no,
                slide_no=slide_no,
                sheet_name=sheet_name,
                row_start=row_start,
                row_end=row_end,
                section_title=section_title,
                metadata=metadata or {},
            )
        )


def safe_pdf_meta_text(reader: PdfReader) -> str:
    try:
        metadata = reader.metadata or {}
    except Exception:
        return ""
    values = []
    for value in metadata.values():
        if value is not None:
            values.append(str(value))
    return " ".join(values).lower()


def pdf_page_texts(reader: PdfReader) -> list[str]:
    texts: list[str] = []
    for page in reader.pages:
        texts.append((page.extract_text() or "").strip())
    return texts


def page_is_landscape(page: Any) -> bool:
    box = page.mediabox
    width = float(box.width)
    height = float(box.height)
    rotation = int(page.get("/Rotate") or 0) % 180
    if rotation == 90:
        width, height = height, width
    return width > height


def classify_pdf(reader: PdfReader, texts: list[str]) -> PdfProfile:
    page_count = max(1, len(reader.pages))
    text_len = sum(len(text) for text in texts)
    avg_chars_per_page = text_len / page_count
    metadata_text = safe_pdf_meta_text(reader)
    metadata_hits = ("powerpoint", "presentation", "keynote", "wps presentation", "slides")

    landscape_count = 0
    for page in reader.pages:
        if page_is_landscape(page):
            landscape_count += 1
    landscape_ratio = landscape_count / page_count

    lines = [line.strip() for text in texts for line in text.splitlines() if line.strip()]
    avg_line_len = sum(len(line) for line in lines) / len(lines) if lines else 0
    short_line_ratio = (
        sum(1 for line in lines if len(line) <= 28) / len(lines)
        if lines
        else 0
    )

    if avg_chars_per_page < TEXT_DENSITY_MIN_CHARS_PER_PAGE:
        kind = "scanned_pdf"
        slide_score = 0
    else:
        slide_score = 0
        if any(hit in metadata_text for hit in metadata_hits):
            slide_score += 3
        if landscape_ratio >= 0.6:
            slide_score += 2
        if short_line_ratio >= 0.55 and len(lines) >= 10:
            slide_score += 1
        if avg_line_len and avg_line_len <= 32:
            slide_score += 1
        if 60 <= avg_chars_per_page <= 700:
            slide_score += 1
        kind = "slide_pdf" if slide_score >= SLIDE_SCORE_THRESHOLD else "document_pdf"

    return PdfProfile(
        kind=kind,
        page_count=page_count,
        text_len=text_len,
        avg_chars_per_page=avg_chars_per_page,
        landscape_ratio=landscape_ratio,
        short_line_ratio=short_line_ratio,
        avg_line_len=avg_line_len,
        slide_score=slide_score,
        metadata={
            "pdf_profile": kind,
            "pdf_text_len": text_len,
            "pdf_avg_chars_per_page": round(avg_chars_per_page, 2),
            "pdf_landscape_ratio": round(landscape_ratio, 2),
            "pdf_short_line_ratio": round(short_line_ratio, 2),
            "pdf_avg_line_len": round(avg_line_len, 2),
            "pdf_slide_score": slide_score,
        },
    )


def iter_docx_blocks(doc: DocxDocument) -> list[Paragraph | Table]:
    blocks: list[Paragraph | Table] = []
    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            blocks.append(Paragraph(child, doc))
        elif child.tag.endswith("}tbl"):
            blocks.append(Table(child, doc))
    return blocks


class ParseProcessor:
    def __init__(self, settings: Settings, storage: ObjectStorage) -> None:
        self.settings = settings
        self.storage = storage
        self.ocr = PaddleOcrClient(settings)

    def process(self, file: FileRow) -> ParseResult:
        ext = normalize_ext(file.file_ext or Path(file.file_name).suffix)
        source = self.storage.get_bytes(file.source_storage_key)

        if ext in TEXT_EXT:
            return self._parse_text(decode_text(source.body), ext)
        if ext in MARKDOWN_EXT:
            return self._parse_markdown(decode_text(source.body), ext)
        if ext in CSV_EXT:
            return self._parse_csv(source.body, "\t" if ext == "tsv" else ",", ext)
        if ext in DOCX_EXT:
            return self._parse_docx(source.body)
        if ext in XLSX_EXT:
            return self._parse_xlsx(source.body)
        if ext in PDF_EXT:
            return self._parse_pdf_or_ocr(source.body, ext, source_kind="pdf")
        if ext in PPT_EXT:
            pdf = self._get_preview_pdf_or_convert(file, source.body)
            return self._parse_ocr_pdf(pdf, source_kind=ext, converted_from=ext)
        if ext in OLD_OFFICE_EXT:
            pdf = self._get_preview_pdf_or_convert(file, source.body)
            return self._parse_pdf_or_ocr(pdf, "pdf", source_kind=ext, converted_from=ext)

        return ParseResult(
            status="skipped",
            output={"version": 1, "reason": "unsupported_parse_type", "ext": ext},
        )

    def _get_preview_pdf_or_convert(self, file: FileRow, source_body: bytes) -> bytes:
        if file.preview_storage_key:
            return self.storage.get_bytes(file.preview_storage_key).body
        return convert_office_to_pdf(
            self.settings,
            file_name=file.file_name,
            data=source_body,
        )

    def _finalize(
        self,
        *,
        chunks: list[ParsedChunk],
        metadata: dict[str, Any],
        parser_name: str = "iws-file-worker-parser",
    ) -> ParseResult:
        content_text = "\n\n".join(c.content for c in chunks if c.content.strip())
        if not chunks or not content_text.strip():
            return ParseResult(
                status="skipped",
                output={"version": 1, "reason": "no_extractable_text", **metadata},
            )
        for idx, chunk in enumerate(chunks):
            chunk.chunk_index = idx
        return ParseResult(
            status="ready",
            content_text=content_text,
            chunks=chunks,
            parser_name=parser_name,
            metadata={"version": 1, **metadata, "chunk_count": len(chunks)},
            output={"version": 1, **metadata, "chunk_count": len(chunks)},
        )

    def _parse_text(self, text: str, ext: str) -> ParseResult:
        chunks = [
            ParsedChunk(i, content, metadata={"source_type": ext, "parse_method": "text_extract"})
            for i, content in enumerate(split_long_text(text))
        ]
        return self._finalize(chunks=chunks, metadata={"source_type": ext, "parse_method": "text_extract"})

    def _parse_markdown(self, text: str, ext: str) -> ParseResult:
        sections = re.split(r"(?m)(?=^#{1,6}\s+)", text)
        chunks: list[ParsedChunk] = []
        for section in sections:
            title_match = re.match(r"(?m)^#{1,6}\s+(.+)$", section.strip())
            title = title_match.group(1).strip() if title_match else None
            for content in split_long_text(section):
                chunks.append(
                    ParsedChunk(
                        len(chunks),
                        content,
                        section_title=title,
                        metadata={"source_type": ext, "parse_method": "text_extract"},
                    )
                )
        return self._finalize(chunks=chunks, metadata={"source_type": ext, "parse_method": "text_extract"})

    def _parse_csv(self, data: bytes, delimiter: str, ext: str) -> ParseResult:
        rows = list(csv.reader(io.StringIO(decode_text(data)), delimiter=delimiter))
        chunks: list[ParsedChunk] = []
        for start in range(0, len(rows), EXCEL_WINDOW_ROWS):
            window = rows[start : start + EXCEL_WINDOW_ROWS]
            content = "\n".join(row_to_text(row) for row in window).strip()
            if content:
                chunks.append(
                    ParsedChunk(
                        len(chunks),
                        content,
                        row_start=start + 1,
                        row_end=start + len(window),
                        metadata={"source_type": ext, "parse_method": "text_extract"},
                    )
                )
        return self._finalize(chunks=chunks, metadata={"source_type": ext, "parse_method": "text_extract"})

    def _parse_docx(self, data: bytes) -> ParseResult:
        doc = Document(io.BytesIO(data))
        chunks: list[ParsedChunk] = []
        current_title: str | None = None
        section_lines: list[str] = []
        table_index = 0
        metadata = {"source_type": "docx", "parse_method": "text_extract"}

        def flush_section() -> None:
            if not section_lines:
                return
            append_text_chunks(
                chunks,
                "\n\n".join(section_lines),
                section_title=current_title,
                metadata=metadata,
            )
            section_lines.clear()

        for block in iter_docx_blocks(doc):
            if isinstance(block, Paragraph):
                text = block.text.strip()
                if not text:
                    continue
                style = (block.style.name or "").lower() if block.style else ""
                if style.startswith("heading"):
                    flush_section()
                    current_title = text
                    section_lines.append(text)
                    continue
                section_lines.append(text)
                continue

            flush_section()
            table_index += 1
            lines = []
            for row in block.rows:
                lines.append(row_to_text([cell.text.strip() for cell in row.cells]))
            append_text_chunks(
                chunks,
                "\n".join(lines).strip(),
                section_title=current_title or f"表格 {table_index}",
                metadata={
                    "source_type": "docx",
                    "parse_method": "text_extract",
                    "block_type": "table",
                    "table_index": table_index,
                },
            )
        flush_section()
        return self._finalize(chunks=chunks, metadata={"source_type": "docx", "parse_method": "text_extract"})

    def _parse_xlsx(self, data: bytes) -> ParseResult:
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        chunks: list[ParsedChunk] = []
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            for start in range(0, len(rows), EXCEL_WINDOW_ROWS):
                window = rows[start : start + EXCEL_WINDOW_ROWS]
                content = "\n".join(row_to_text(list(row)) for row in window).strip()
                if not content:
                    continue
                chunks.append(
                    ParsedChunk(
                        len(chunks),
                        content,
                        sheet_name=sheet.title,
                        row_start=start + 1,
                        row_end=start + len(window),
                        metadata={"source_type": "xlsx", "parse_method": "text_extract"},
                    )
                )
        return self._finalize(chunks=chunks, metadata={"source_type": "xlsx", "parse_method": "text_extract"})

    def _parse_pdf_or_ocr(
        self,
        data: bytes,
        ext: str,
        *,
        source_kind: str,
        converted_from: str | None = None,
    ) -> ParseResult:
        reader = PdfReader(io.BytesIO(data))
        texts = pdf_page_texts(reader)
        profile = classify_pdf(reader, texts)
        if profile.kind == "document_pdf":
            return self._parse_pdf_text(
                reader,
                texts,
                profile,
                source_kind=source_kind,
                converted_from=converted_from,
            )
        return self._parse_ocr_pdf(
            data,
            source_kind=source_kind,
            converted_from=converted_from,
            profile_metadata=profile.metadata,
        )

    def _parse_pdf_text(
        self,
        reader: PdfReader,
        texts: list[str],
        profile: PdfProfile,
        *,
        source_kind: str,
        converted_from: str | None,
    ) -> ParseResult:
        chunks: list[ParsedChunk] = []
        page_buffer: list[str] = []
        page_start: int | None = None
        page_end: int | None = None

        def flush_pages() -> None:
            nonlocal page_start, page_end
            if not page_buffer or page_start is None or page_end is None:
                return
            metadata = {
                "source_type": source_kind,
                "parse_method": "text_extract",
                "converted_from": converted_from,
                **profile.metadata,
                "page_start": page_start,
                "page_end": page_end,
            }
            for content in split_long_text("\n\n".join(page_buffer)):
                chunks.append(
                    ParsedChunk(
                        len(chunks),
                        content,
                        page_no=page_start,
                        metadata=metadata,
                    )
                )
            page_buffer.clear()
            page_start = None
            page_end = None

        for page_idx, text in enumerate(texts, start=1):
            if not text:
                continue
            if len(text) > MAX_CHUNK_CHARS:
                flush_pages()
                metadata = {
                    "source_type": source_kind,
                    "parse_method": "text_extract",
                    "converted_from": converted_from,
                    **profile.metadata,
                    "page_start": page_idx,
                    "page_end": page_idx,
                }
                for content in split_long_text(text):
                    chunks.append(
                        ParsedChunk(
                            len(chunks),
                            content,
                            page_no=page_idx,
                            metadata=metadata,
                        )
                    )
                continue

            buffer_text_len = len("\n\n".join(page_buffer))
            if page_buffer and buffer_text_len + len(text) + 2 > MAX_CHUNK_CHARS:
                flush_pages()

            if page_start is None:
                page_start = page_idx
            page_end = page_idx
            page_buffer.append(text)

        flush_pages()
        return self._finalize(
            chunks=chunks,
            metadata={
                "source_type": source_kind,
                "parse_method": "text_extract",
                "converted_from": converted_from,
                "page_count": profile.page_count,
                **profile.metadata,
            },
        )

    def _parse_ocr_pdf(
        self,
        data: bytes,
        *,
        source_kind: str,
        converted_from: str | None = None,
        profile_metadata: dict[str, Any] | None = None,
    ) -> ParseResult:
        pages = self.ocr.parse(data, "pdf")
        base_metadata = profile_metadata or {}
        chunks: list[ParsedChunk] = []
        for page in pages:
            for content in split_long_text(page.text):
                chunks.append(
                    ParsedChunk(
                        len(chunks),
                        content,
                        page_no=page.page_no,
                        slide_no=page.page_no if source_kind in PPT_EXT else None,
                        metadata={
                            "source_type": source_kind,
                            "parse_method": "ocr",
                            "ocr_provider": "paddleocr",
                            "converted_from": converted_from,
                            **base_metadata,
                        },
                    )
                )
        return self._finalize(
            chunks=chunks,
            metadata={
                "source_type": source_kind,
                "parse_method": "ocr",
                "ocr_provider": "paddleocr",
                "converted_from": converted_from,
                "page_count": len(pages),
                **base_metadata,
            },
            parser_name="iws-file-worker-paddleocr",
        )
